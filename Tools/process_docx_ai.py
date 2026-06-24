import os
import json
import sys
import docx
import requests

API_KEY = "sk-92d3ef3c86e600b6-ef9847-d92a108c"
API_URL = "http://localhost:20128/v1/chat/completions"
MODEL_ID = "tllm/GPT_5_4"

SYSTEM_PROMPT = """You are an expert IELTS data processor.
Your task is to take messy, raw text extracted from an IELTS test document (Listening or Reading) and convert it into a perfectly structured JSON object.

RULES:
1. Ignore any garbage headers or footers.
2. Convert any blank lines (like '___' or '...') into the standard placeholder `{{BLANK}}`.
3. Separate the content logically into parts/sections.
4. If there are passages (for reading or context for listening), extract them cleanly.
5. If there are questions, identify the question type (e.g., "fill_in_blank", "multiple_choice", "matching") and extract the question numbers, text, and options.
6. The output MUST be a valid JSON object. Do not include markdown formatting like ```json at the beginning or end. Just return the JSON.

RETURN ONLY VALID JSON.

JSON SCHEMA:
{
  "test_title": "string",
  "sections": [
    {
      "section_name": "string (e.g., Part 1: Questions 1-10)",
      "instruction": "string",
      "passage_or_context": "string (optional)",
      "questions": [
        {
          "q_numbers": "string (e.g., 1-5 or 11)",
          "type": "string (fill_in_blank, multiple_choice, etc.)",
          "content": "string (The actual question text or the sentence with {{BLANK}} placeholders)",
          "options": ["string"] // optional, for multiple choice or matching
        }
      ]
    }
  ]
}
"""

def extract_text_from_docx(file_path):
    doc = docx.Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())
    return "\n".join(full_text)

def process_with_ai(text, model_id=MODEL_ID):
    headers = {
        "Authorization": f"Bearer {API_KEY}",
        "Content-Type": "application/json"
    }
    
    payload = {
        "model": model_id,
        "stream": False,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": text}
        ]
    }
    
    response = requests.post(API_URL, headers=headers, json=payload, verify=False)
    if response.status_code == 200:
        try:
            data = response.json()
            result_text = data['choices'][0]['message']['content']
            # clean up backticks just in case
            if result_text.startswith('```json'):
                result_text = result_text[7:]
            if result_text.startswith('```'):
                result_text = result_text[3:]
            if result_text.endswith('```'):
                result_text = result_text[:-3]
            return result_text.strip()
        except Exception as e:
            raise Exception(f"Failed to parse JSON. Raw response: {response.text[:500]}")
    else:
        raise Exception(f"API Error: {response.status_code} - {response.text}")

def main():
    if len(sys.argv) < 2:
        print("Usage: python process_docx_ai.py <path_to_docx>")
        sys.exit(1)
        
    file_path = sys.argv[1]
    print(f"Reading {file_path}...")
    raw_text = extract_text_from_docx(file_path)
    
    print("Sending to OmniRoute API...")
    try:
        json_output = process_with_ai(raw_text)
        
        # Sometimes the AI still includes markdown code blocks
        if json_output.startswith("```json"):
            json_output = json_output[7:]
        if json_output.endswith("```"):
            json_output = json_output[:-3]
            
        json_output = json_output.strip()
        
        # Verify it's valid JSON
        parsed_json = json.loads(json_output)
        
        output_path = file_path.replace('.docx', '_processed.json')
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(parsed_json, f, ensure_ascii=False, indent=2)
            
        print(f"Success! Processed JSON saved to {output_path}")
        
    except Exception as e:
        print(f"Failed to process {file_path}: {e}")

if __name__ == "__main__":
    main()
